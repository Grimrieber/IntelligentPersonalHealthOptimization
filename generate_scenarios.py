"""
Generate Excel spreadsheet with all nutrition formula scenarios
using Stephan's personal data against every goal/timeline/diet combination.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import math

# ===== Stephan's Profile =====
WEIGHT_KG = 100.7
HEIGHT_CM = 178.2
AGE = 40
GENDER = "Male"
ACTIVITY_LEVEL = "ModeratelyActive"
ACTIVITY_FACTOR = 1.55
TARGET_WEIGHT_KG = 91.6  # ~201.9 lb
READINESS_SCORE = 4.7

# ===== BMR & TDEE =====
BMR = 10 * WEIGHT_KG + 6.25 * HEIGHT_CM - 5 * AGE + 5  # Male
TDEE = BMR * ACTIVITY_FACTOR

# ===== Enums =====
GOALS = [
    "LoseWeight", "LoseBodyFat", "BodyRecomposition", "PrepareForCompetition",
    "BuildLeanMuscle", "AggressiveMuscleGain", "ImproveAthletePerformance",
    "ImproveEndurance", "IncreaseEnergy", "MaintainWeight",
    "ImproveOverallHealth", "ManageHealthCondition"
]

GOAL_DEFAULTS = {
    "LoseWeight": -500, "LoseBodyFat": -500, "BodyRecomposition": -250,
    "PrepareForCompetition": -400, "BuildLeanMuscle": 250,
    "AggressiveMuscleGain": 500, "ImproveAthletePerformance": 200,
    "ImproveEndurance": 150, "IncreaseEnergy": 150, "MaintainWeight": 0,
    "ImproveOverallHealth": 0, "ManageHealthCondition": 0
}

GOAL_MACROS = {
    "LoseWeight": (0.35, 0.35, 0.30),
    "LoseBodyFat": (0.40, 0.30, 0.30),
    "BodyRecomposition": (0.40, 0.35, 0.25),
    "PrepareForCompetition": (0.45, 0.25, 0.30),
    "BuildLeanMuscle": (0.30, 0.45, 0.25),
    "AggressiveMuscleGain": (0.25, 0.50, 0.25),
    "ImproveAthletePerformance": (0.25, 0.50, 0.25),
    "ImproveEndurance": (0.20, 0.55, 0.25),
    "IncreaseEnergy": (0.25, 0.50, 0.25),
    "MaintainWeight": (0.30, 0.40, 0.30),
    "ImproveOverallHealth": (0.30, 0.40, 0.30),
    "ManageHealthCondition": (0.30, 0.40, 0.30),
}

TIMELINES = {
    "SixWeeks": (42, 1.3),
    "EightWeeks": (56, 1.2),
    "TwelveWeeks": (84, 1.0),
    "SixMonths": (182, 0.7),
}

DIET_TYPES = ["Standard", "Keto", "Paleo", "Mediterranean", "Vegetarian", "Vegan"]

DIET_MACRO_OVERRIDES = {
    "Keto": (0.25, 0.05, 0.70),
    "Paleo": (0.30, 0.25, 0.45),
    "Mediterranean": (0.20, 0.45, 0.35),
}

ACTIVITY_LEVELS = {
    "Sedentary": 1.2,
    "LightlyActive": 1.375,
    "ModeratelyActive": 1.55,
    "VeryActive": 1.725,
    "ExtremelyActive": 1.9,
}


def calc_bmr(weight_kg, height_cm, age, gender):
    bmr = 10 * weight_kg + 6.25 * height_cm - 5 * age
    return bmr + 5 if gender == "Male" else bmr - 161


def calc_targets(tdee, goal, diet, gender, current_wt, target_wt, timeline_key, readiness):
    timeline_days, timeline_intensity = TIMELINES[timeline_key]

    # Weight-based adjustment
    if target_wt > 0 and current_wt > 0 and abs(target_wt - current_wt) > 0.5:
        weight_diff = target_wt - current_wt
        weight_adjust = weight_diff * 7700.0 / timeline_days
        daily_adj = max(-1000, min(750, weight_adjust * timeline_intensity))
    else:
        daily_adj = GOAL_DEFAULTS.get(goal, 0) * timeline_intensity

    # Readiness dampening
    if readiness < 4.0:
        dampen = 0.5 + (readiness / 4.0) * 0.5
        daily_adj *= dampen

    # Safety floor
    cal_floor = 1500 if gender == "Male" else 1200
    target_cal = max(tdee + daily_adj, cal_floor)

    # Macro split
    if diet in DIET_MACRO_OVERRIDES:
        p_pct, c_pct, f_pct = DIET_MACRO_OVERRIDES[diet]
    else:
        p_pct, c_pct, f_pct = GOAL_MACROS.get(goal, (0.30, 0.40, 0.30))

    protein_g = int(target_cal * p_pct / 4)
    carbs_g = int(target_cal * c_pct / 4)
    fat_g = int(target_cal * f_pct / 9)

    return int(target_cal), daily_adj, protein_g, carbs_g, fat_g, p_pct, c_pct, f_pct


# ===== Styling =====
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="5B21B6", end_color="5B21B6", fill_type="solid")
info_fill = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
deficit_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
surplus_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
maint_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB'),
)


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = thin_border


def style_cell(cell, fill=None):
    cell.border = thin_border
    cell.alignment = Alignment(horizontal='center')
    if fill:
        cell.fill = fill


wb = openpyxl.Workbook()

# ==================== SHEET 1: Profile Summary ====================
ws = wb.active
ws.title = "Profile"
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 30

profile_data = [
    ("Your Profile", ""),
    ("Name", "Stephan"),
    ("Weight", f"{WEIGHT_KG} kg ({WEIGHT_KG * 2.20462:.0f} lb)"),
    ("Height", f"{HEIGHT_CM} cm"),
    ("Age", AGE),
    ("Gender", GENDER),
    ("Activity Level", f"{ACTIVITY_LEVEL} ({ACTIVITY_FACTOR}x)"),
    ("Target Weight", f"{TARGET_WEIGHT_KG} kg ({TARGET_WEIGHT_KG * 2.20462:.0f} lb)"),
    ("Weight to Change", f"{TARGET_WEIGHT_KG - WEIGHT_KG:.1f} kg ({(TARGET_WEIGHT_KG - WEIGHT_KG) * 2.20462:.1f} lb)"),
    ("Readiness Score", f"{READINESS_SCORE}/10"),
    ("", ""),
    ("Calculated Values", ""),
    ("BMR (Mifflin-St Jeor)", f"{BMR:.0f} kcal"),
    ("TDEE", f"{TDEE:.0f} kcal"),
    ("", ""),
    ("Formula", ""),
    ("BMR", "10 × weight(kg) + 6.25 × height(cm) - 5 × age + 5(M)/-161(F)"),
    ("TDEE", "BMR × Activity Multiplier"),
    ("Weight Adjust", "weightDiff(kg) × 7700 / timelineDays × timelineIntensity"),
    ("Calories", "max(TDEE + dailyAdjustment, calorieFloor)"),
    ("Protein", "calories × proteinPct / 4"),
    ("Carbs", "calories × carbsPct / 4"),
    ("Fat", "calories × fatPct / 9"),
    ("Calorie Floor", "Male: 1500 | Female: 1200"),
]

for i, (label, value) in enumerate(profile_data, 1):
    ws.cell(row=i, column=1, value=label).font = Font(bold=True) if label else Font()
    ws.cell(row=i, column=2, value=value)
    if i == 1 or label in ("Calculated Values", "Formula"):
        ws.cell(row=i, column=1).fill = info_fill
        ws.cell(row=i, column=2).fill = info_fill

# ==================== SHEET 2: All Goals × Timelines (Standard Diet) ====================
ws2 = wb.create_sheet("Goals × Timelines")

headers = [
    "Goal", "Timeline", "Timeline Days", "Intensity",
    "Weight Diff (kg)", "Daily Adjust", "TDEE", "Target Calories",
    "Deficit/Surplus", "Protein %", "Carbs %", "Fat %",
    "Protein (g)", "Carbs (g)", "Fat (g)"
]

for c, h in enumerate(headers, 1):
    ws2.cell(row=1, column=c, value=h)
style_header(ws2, 1, len(headers))

row = 2
for goal in GOALS:
    for tl_name, (tl_days, tl_intensity) in TIMELINES.items():
        cal, adj, p, c_g, f, pp, cp, fp = calc_targets(
            TDEE, goal, "Standard", GENDER, WEIGHT_KG, TARGET_WEIGHT_KG, tl_name, READINESS_SCORE)

        values = [
            goal, tl_name, tl_days, tl_intensity,
            round(TARGET_WEIGHT_KG - WEIGHT_KG, 1), round(adj, 0), round(TDEE, 0), cal,
            f"{'+' if adj >= 0 else ''}{int(adj)}", f"{pp:.0%}", f"{cp:.0%}", f"{fp:.0%}",
            p, c_g, f
        ]

        fill = deficit_fill if adj < -10 else surplus_fill if adj > 10 else maint_fill
        for ci, v in enumerate(values, 1):
            cell = ws2.cell(row=row, column=ci, value=v)
            style_cell(cell, fill)

        row += 1

for c in range(1, len(headers) + 1):
    ws2.column_dimensions[get_column_letter(c)].width = 16
ws2.column_dimensions['A'].width = 28

# ==================== SHEET 3: All Goals × Diet Types (12 weeks) ====================
ws3 = wb.create_sheet("Goals × Diets")

headers3 = [
    "Goal", "Diet Type", "Target Calories", "Daily Adjust",
    "Protein %", "Carbs %", "Fat %",
    "Protein (g)", "Carbs (g)", "Fat (g)"
]

for c, h in enumerate(headers3, 1):
    ws3.cell(row=1, column=c, value=h)
style_header(ws3, 1, len(headers3))

row = 2
for goal in GOALS:
    for diet in DIET_TYPES:
        cal, adj, p, c_g, f, pp, cp, fp = calc_targets(
            TDEE, goal, diet, GENDER, WEIGHT_KG, TARGET_WEIGHT_KG, "TwelveWeeks", READINESS_SCORE)

        values = [
            goal, diet, cal, f"{'+' if adj >= 0 else ''}{int(adj)}",
            f"{pp:.0%}", f"{cp:.0%}", f"{fp:.0%}",
            p, c_g, f
        ]

        fill = deficit_fill if adj < -10 else surplus_fill if adj > 10 else maint_fill
        for ci, v in enumerate(values, 1):
            cell = ws3.cell(row=row, column=ci, value=v)
            style_cell(cell, fill)

        row += 1

for c in range(1, len(headers3) + 1):
    ws3.column_dimensions[get_column_letter(c)].width = 16
ws3.column_dimensions['A'].width = 28

# ==================== SHEET 4: Activity Level Impact ====================
ws4 = wb.create_sheet("Activity Levels")

headers4 = [
    "Activity Level", "Multiplier", "BMR", "TDEE",
    "LoseWeight Cal", "LoseWeight Deficit",
    "BuildLeanMuscle Cal", "BuildLeanMuscle Adjust",
    "MaintainWeight Cal"
]

for c, h in enumerate(headers4, 1):
    ws4.cell(row=1, column=c, value=h)
style_header(ws4, 1, len(headers4))

row = 2
for al_name, al_factor in ACTIVITY_LEVELS.items():
    al_tdee = BMR * al_factor

    lw_cal, lw_adj, _, _, _, _, _, _ = calc_targets(
        al_tdee, "LoseWeight", "Standard", GENDER, WEIGHT_KG, TARGET_WEIGHT_KG, "TwelveWeeks", READINESS_SCORE)
    blm_cal, blm_adj, _, _, _, _, _, _ = calc_targets(
        al_tdee, "BuildLeanMuscle", "Standard", GENDER, WEIGHT_KG, TARGET_WEIGHT_KG, "TwelveWeeks", READINESS_SCORE)
    mw_cal, _, _, _, _, _, _, _ = calc_targets(
        al_tdee, "MaintainWeight", "Standard", GENDER, WEIGHT_KG, TARGET_WEIGHT_KG, "TwelveWeeks", READINESS_SCORE)

    values = [
        al_name, al_factor, round(BMR, 0), round(al_tdee, 0),
        lw_cal, int(lw_adj), blm_cal, int(blm_adj), mw_cal
    ]

    for ci, v in enumerate(values, 1):
        cell = ws4.cell(row=row, column=ci, value=v)
        style_cell(cell)

    row += 1

for c in range(1, len(headers4) + 1):
    ws4.column_dimensions[get_column_letter(c)].width = 20
ws4.column_dimensions['A'].width = 22

# ==================== SHEET 5: No Target Weight (goal defaults) ====================
ws5 = wb.create_sheet("No Target Weight")

headers5 = [
    "Goal", "Default Adjust", "12-Week Cal", "6-Week Cal", "6-Month Cal",
    "Protein (g)", "Carbs (g)", "Fat (g)"
]

for c, h in enumerate(headers5, 1):
    ws5.cell(row=1, column=c, value=h)
style_header(ws5, 1, len(headers5))

row = 2
for goal in GOALS:
    cal_12, adj_12, p, c_g, f, _, _, _ = calc_targets(
        TDEE, goal, "Standard", GENDER, WEIGHT_KG, 0, "TwelveWeeks", READINESS_SCORE)
    cal_6w, _, _, _, _, _, _, _ = calc_targets(
        TDEE, goal, "Standard", GENDER, WEIGHT_KG, 0, "SixWeeks", READINESS_SCORE)
    cal_6m, _, _, _, _, _, _, _ = calc_targets(
        TDEE, goal, "Standard", GENDER, WEIGHT_KG, 0, "SixMonths", READINESS_SCORE)

    values = [
        goal, GOAL_DEFAULTS.get(goal, 0), cal_12, cal_6w, cal_6m, p, c_g, f
    ]

    fill = deficit_fill if GOAL_DEFAULTS.get(goal, 0) < 0 else surplus_fill if GOAL_DEFAULTS.get(goal, 0) > 0 else maint_fill
    for ci, v in enumerate(values, 1):
        cell = ws5.cell(row=row, column=ci, value=v)
        style_cell(cell, fill)

    row += 1

for c in range(1, len(headers5) + 1):
    ws5.column_dimensions[get_column_letter(c)].width = 18
ws5.column_dimensions['A'].width = 28

# ===== Save =====
output_path = r"D:\Claude Projects\IntelligentPersonalHealthOptimization\NutritionFormulaScenarios.xlsx"
wb.save(output_path)
print(f"Saved to {output_path}")
